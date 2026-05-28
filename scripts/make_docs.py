# -*- coding: utf-8 -*-
"""Generate Buleto system documentation as a formatted .docx (admin + customer flows)."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ACCENT = RGBColor(0xFF, 0x50, 0x46)
INK = RGBColor(0x1F, 0x23, 0x29)
GREY = RGBColor(0x6C, 0x75, 0x7D)

doc = Document()

# ---- base styles ----
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.font.color.rgb = INK

for lvl, size, color in [('Heading 1', 17, ACCENT), ('Heading 2', 14, INK), ('Heading 3', 12, INK)]:
    st = doc.styles[lvl]
    st.font.name = 'Calibri'
    st.font.size = Pt(size)
    st.font.color.rgb = color
    st.font.bold = True


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:fill'), hexcolor)
    tcPr.append(sh)


def para(text='', bold=False, italic=False, color=None, size=None, align=None, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = color
    if size:
        r.font.size = Pt(size)
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p


def number(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Number')
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ''
        run = hdr[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(10)
        shade(hdr[i], '1F2329')
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            r = cells[i].paragraphs[0].add_run(str(val))
            r.font.size = Pt(10)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t


def hr():
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'FF5046')
    pbdr.append(bottom)
    pPr.append(pbdr)


# ============================ COVER ============================
for _ in range(4):
    doc.add_paragraph()
para('BULETO', bold=True, color=ACCENT, size=40, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para('Online Ticket Booking Platform', bold=True, size=20, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para('System Documentation — Admin & Customer Flows', size=13, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
para('Movies • Events • Sports ticketing  |  Laravel 13 + MySQL', size=11, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)
para('Version 1.0', size=11, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

# ============================ 1. OVERVIEW ============================
doc.add_heading('1. System Overview', level=1)
para('Buleto is an online ticket-booking platform (BookMyShow style) for cinemas, events and sports. '
     'It is a single Laravel application with a public website for customers and a separate admin panel for staff. '
     'Customers browse what is on, pick a showtime, choose seats on an interactive seat map, pay, and receive a '
     'QR-coded e-ticket. Admins manage the entire catalogue, scheduling, pricing and content from the dashboard.')

doc.add_heading('Technology', level=3)
table(['Layer', 'Technology'], [
    ['Framework', 'Laravel 13 (PHP 8.3)'],
    ['Database', 'MySQL 8.4'],
    ['Frontend', 'Blade templates + Bootstrap (admin), themed HTML/JS (public site)'],
    ['Seat locking', 'Laravel atomic cache locks (database store) — equivalent to Redis NX'],
    ['Media', 'UniSharp Laravel File Manager'],
    ['Payments', 'eSewa (sandbox), Khalti, Card/Mock'],
    ['Email', 'Laravel Mail (log driver in dev; SendGrid/SES ready)'],
], widths=[1.8, 4.5])

doc.add_heading('Two roles, two entrances', level=3)
table(['', 'Customer', 'Administrator'], [
    ['Login URL', '/login', '/admin/login (separate, branded page)'],
    ['Home after login', '/account (My Account)', '/admin (Dashboard)'],
    ['Test credentials', 'user@buleto.test / password', 'admin@buleto.test / password'],
    ['Can register?', 'Yes, at /register', 'No — created by another admin'],
], widths=[1.4, 2.6, 2.8])
para('Note: the two logins are completely separate. A customer who tries the admin login is rejected; '
     'unauthenticated visits to any /admin page are redirected to the admin login.', italic=True, color=GREY, size=10)

doc.add_page_break()

# ============================ 2. ADMIN FLOW ============================
doc.add_heading('2. Administrator Flow', level=1)

doc.add_heading('2.1 Logging in', level=2)
number('Go to ', bold_prefix='/admin/login')
number('Enter the admin email and password (admin@buleto.test / password for the test account).')
number('On success you land on the Admin Dashboard. The left sidebar is your main navigation.')
para('The admin panel is organised into sections in the sidebar: Overview, Movies, Events, Sports, Blog, Site, and Media.',
     space_after=4)

doc.add_heading('2.2 The Dashboard', level=2)
para('The dashboard shows summary statistics (totals for movies, bookings, users, revenue, etc.) so staff can see '
     'the state of the business at a glance. Use the sidebar to jump to any management area.')

doc.add_heading('2.3 How adding & editing works (the universal pattern)', level=2)
para('Almost every management screen follows the same Create / Read / Update / Delete (CRUD) pattern, so once you '
     'learn one, you know them all:')
number('Click a section in the sidebar (e.g. Movies). You see a LIST of existing records with search and pagination.', bold_prefix='List: ')
number('Click "Create" (top-right) to add a new record. Fill the form and click Save.', bold_prefix='Create: ')
number('Click a row\'s "Edit" to change it; click "Delete" to remove it (with confirmation).', bold_prefix='Edit/Delete: ')
para('The forms are generated automatically from each record\'s fields:', space_after=2)
bullet('Text/number/date fields appear as the matching input type.')
bullet('A field ending in "_id" (e.g. screen_id) becomes a dropdown of related records, shown by name.')
bullet('Image fields (poster, banner, etc.) show a "Choose" button that opens the File Manager to pick or upload an image.')
para('Example — the Showtime "Screen" dropdown shows "Star Cinema — Screen 2 (#8)" so you always know which '
     'cinema a screen belongs to.', italic=True, color=GREY, size=10)

doc.add_heading('2.4 What you can manage', level=2)
table(['Section', 'What it controls', 'Where it shows for customers'], [
    ['Movies', 'Title, synopsis, poster, languages, genres, formats, cast, gallery', 'Movies listing & detail pages'],
    ['Cinemas', 'Venue name, city, address, location', 'Shown on showtime & seat pages'],
    ['Screens', 'A hall inside a cinema + its seat layout (rows × seats)', 'Drives the seat map grid'],
    ['Showtimes', 'A movie playing on a screen at a date/time, with language & format', 'Showtimes list → booking'],
    ['Ticket Classes', 'Seat price tiers (e.g. Silver/Gold/VIP) per showtime', 'Price selector on seat page'],
    ['Languages / Formats / Genres', 'Reference lists used by movies & showtimes', 'Filters & labels'],
    ['Promo Codes', 'Discount codes', 'Checkout promo field'],
    ['Popcorn Items', 'Food & beverage add-ons', 'Add-ons / popcorn page'],
    ['Events / Event Categories / Speakers', 'Conferences, concerts, etc.', 'Events section'],
    ['Sports / Sport Categories', 'Matches & sport ticketing', 'Sports section'],
    ['Blog Posts / Categories / Tags / Comments', 'Articles and moderation', 'Blog section'],
    ['Users', 'Customer & admin accounts and roles', '—'],
    ['Bookings', 'View, inspect and refund customer bookings', 'Reflected in customer account'],
    ['Cities / Banners / FAQs / Partners', 'Homepage & site content', 'Homepage and info pages'],
    ['Settings', 'Global site settings', 'Site-wide'],
    ['File Manager (Media)', 'Upload & organise all images', 'Picked into any image field'],
], widths=[1.7, 3.0, 2.0])

doc.add_heading('2.5 The File Manager (images)', level=2)
para('Open Media → File Manager (or click "Choose" on any image field). You can browse folders (about, banner, '
     'movie, event, etc.), upload new images, crop/resize, and rename. When picked from an image field, the chosen '
     'image\'s path is filled in automatically and a preview appears. Uploaded photos go to a shared media library '
     'reusable across the whole site.')

doc.add_heading('2.6 Seat layout (how the seat map is defined)', level=2)
para('Each Screen has a Seat Layout: a set of rows (A, B, C…) and how many seats are in each row. This is edited '
     'on the Screen form ("Add Row", set seats per row). The customer-facing seat map is generated directly from '
     'this layout — change a screen\'s rows and every showtime on that screen reflects it instantly. "Total Seats" '
     'is auto-calculated from the rows.')

doc.add_heading('2.7 Recipe: publish a movie customers can book', level=2)
para('This is the end-to-end order in which records must exist, because each depends on the previous one:')
number('Create reference data if missing: Languages, Formats, Genres, a City.', bold_prefix='Reference: ')
number('Create the Cinema (venue) in that city.', bold_prefix='Cinema: ')
number('Create one or more Screens in that cinema, each with a seat layout.', bold_prefix='Screen: ')
number('Create the Movie (title, poster via File Manager, languages/genres/formats).', bold_prefix='Movie: ')
number('Create a Showtime linking the movie + screen + language + format + date + time + available seats.', bold_prefix='Showtime: ')
number('Create Ticket Classes (price tiers) for that showtime.', bold_prefix='Pricing: ')
para('The movie now appears on the site, the showtime is bookable, and the seat map uses the screen\'s layout '
     'with the ticket-class prices. ', space_after=10)

doc.add_page_break()

# ============================ 3. CUSTOMER FLOW ============================
doc.add_heading('3. Customer Flow', level=1)

doc.add_heading('3.1 Discover', level=2)
bullet('Visit the homepage to see featured banners, "Now Showing" movies, events and sports.', bold_prefix='Home: ')
bullet('Browse /movies (with filters), open a movie to read details and see its showtimes.', bold_prefix='Browse: ')

doc.add_heading('3.2 Register / Log in', level=2)
para('Anyone can browse, but to book you must be logged in. New customers sign up at /register (created as a '
     '"customer" role). Existing customers log in at /login. Login is rate-limited (5 attempts per 10 minutes) to '
     'prevent abuse.')

doc.add_heading('3.3 Choose a showtime & seats', level=2)
number('From a movie, open Showtimes and pick a date/time at a cinema.', bold_prefix='Pick showtime: ')
number('The interactive Seat Map opens, showing the screen and all seats grouped by row.', bold_prefix='Seat map: ')
number('Choose a Ticket Class (price tier) and click available seats (up to 10).', bold_prefix='Select: ')
para('Seat colours update live (the map refreshes every 10 seconds):', space_after=2)
table(['Colour', 'Meaning'], [
    ['Grey', 'Available'],
    ['Green', 'Selected by you'],
    ['Yellow', 'Locked — someone else is choosing it right now'],
    ['Red', 'Booked (already sold)'],
], widths=[1.4, 4.5])

doc.add_heading('3.4 Seat locking (no double-booking)', level=2)
para('When you proceed, the seats you picked are LOCKED for 5 minutes so no one else can take them while you pay. '
     'If two people click the same seat at the same instant, only one wins — the other is told the seat was just '
     'taken. If you abandon checkout, the lock expires and the seats are released automatically.')

doc.add_heading('3.5 Checkout & payment', level=2)
number('Review your booking summary (movie, cinema/screen, time, seats, price). A 5-minute hold applies.', bold_prefix='Review: ')
number('Optionally enter a promo code.', bold_prefix='Promo: ')
number('Choose a payment method and confirm.', bold_prefix='Pay: ')
para('Available payment methods (all free to test):', space_after=2)
table(['Method', 'How it behaves'], [
    ['eSewa (sandbox)', 'Redirects to the real eSewa test page; log in with test ID 9806800001 / Nepal@123 / MPIN 1122 / OTP 123456'],
    ['Khalti', 'Uses a mock unless a free Khalti test key is configured in the system'],
    ['Card (test)', 'Auto-completes instantly — best for quick testing'],
], widths=[1.6, 4.7])

doc.add_heading('3.6 Confirmation & e-ticket', level=2)
para('After successful payment the booking is confirmed and you are taken to your e-ticket, which shows the movie, '
     'cinema, screen, date/time, seat numbers, amount paid and a QR code to scan at entry. A confirmation email '
     '(with the QR) is also sent, and an SMS notification is queued.')

doc.add_heading('3.7 My Account', level=2)
bullet('See all your bookings and open any e-ticket again.', bold_prefix='Bookings: ')
bullet('Cancel an eligible booking (frees the seats).', bold_prefix='Cancel: ')
bullet('Update your profile and password.', bold_prefix='Profile: ')

doc.add_page_break()

# ============================ 4. HOW THE SYSTEM WORKS ============================
doc.add_heading('4. How the System Works (Technical)', level=1)

doc.add_heading('4.1 Booking integrity — three layers against double-booking', level=2)
number('Atomic seat locks (5-min TTL) stop two users grabbing a seat during checkout. Locks are all-or-nothing: '
       'if any one seat in your selection is taken, none are locked.', bold_prefix='Cache lock: ')
number('A unique database constraint on (showtime, seat row, seat number) makes it physically impossible for two '
       'confirmed bookings to share a seat — the final guarantee.', bold_prefix='DB unique: ')
number('A scheduled cleanup cancels unpaid bookings older than 15 minutes and frees their seats.', bold_prefix='Cleanup: ')

doc.add_heading('4.2 Payment flow', level=2)
number('Customer confirms → the system creates a Payment record and contacts the gateway (initiate).')
number('Customer completes payment on the gateway (or instantly for Card/Mock).')
number('Gateway returns to the system, which verifies the payment server-to-server.')
number('On success: booking becomes "confirmed", a QR code is generated, seat locks are released, and a '
       '"Booking Confirmed" event fires.')

doc.add_heading('4.3 Notifications', level=2)
para('The "Booking Confirmed" event triggers a listener that emails the e-ticket (with QR) and logs an SMS '
     'notification. In development email is written to the log; in production it can use SendGrid/SES, and SMS '
     'can use Twilio or a Nepali SMS provider.')

doc.add_heading('4.4 Seat-status API', level=2)
para('The seat map calls GET /api/showtimes/{id}/seats, which returns every seat grouped by row with a status of '
     'available, locked, mine or booked. The page polls this endpoint every 10 seconds for live updates.')

doc.add_heading('4.5 Rate limiting', level=2)
table(['Action', 'Limit'], [
    ['Login (admin & customer)', '5 attempts / 10 minutes per IP'],
    ['Seat locking', '10 requests / minute per user'],
    ['Payments', '5 requests / minute per user'],
], widths=[2.6, 3.7])

doc.add_heading('4.6 Key data relationships', level=2)
para('City → Cinema → Screen (with seat layout) → Showtime (also links Movie, Language, Format). '
     'A Booking belongs to a Customer and a Showtime, and has many Booking Seats and one or more Payments. '
     'Ticket Classes set the price tiers shown on the seat page.')

doc.add_page_break()

# ============================ 5. URL REFERENCE ============================
doc.add_heading('5. Quick URL Reference', level=1)
doc.add_heading('Customer', level=3)
table(['Page', 'URL'], [
    ['Home', '/'],
    ['Movies (with filters)', '/movies'],
    ['Movie detail', '/movies/{slug}'],
    ['Showtimes for a movie', '/movies/{slug}/showtimes'],
    ['Seat map', '/showtimes/{id}/seats'],
    ['Checkout', '/checkout/movie/{booking}'],
    ['E-ticket', '/bookings/{booking}/ticket'],
    ['My account', '/account'],
    ['Register / Login', '/register  •  /login'],
], widths=[2.6, 3.7])

doc.add_heading('Administrator', level=3)
table(['Page', 'URL'], [
    ['Admin login', '/admin/login'],
    ['Dashboard', '/admin'],
    ['Movies / Cinemas / Screens / Showtimes', '/admin/movies, /admin/cinemas, /admin/screens, /admin/showtimes'],
    ['Ticket Classes / Promo Codes / Popcorn', '/admin/ticket-classes, /admin/promo-codes, /admin/popcorn-items'],
    ['Events / Sports / Blog', '/admin/events, /admin/sports, /admin/blog-posts'],
    ['Users / Bookings / Settings', '/admin/users, /admin/bookings, /admin/settings'],
    ['File Manager', '/admin/filemanager'],
], widths=[2.6, 3.7])

hr()
para('End of document — Buleto System Documentation v1.0', italic=True, color=GREY, size=10,
     align=WD_ALIGN_PARAGRAPH.CENTER)

out = r'C:\laragon\www\Buleto\Buleto-System-Documentation.docx'
doc.save(out)
print('SAVED', out)
